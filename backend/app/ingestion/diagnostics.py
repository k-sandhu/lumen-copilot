"""Native extraction inspection. Measurements never imply complete visual fidelity."""

from __future__ import annotations

import io
import unicodedata
from collections.abc import Iterator
from xml.etree import ElementTree

from app.domain.ingestion import ExtractionDiagnostics, LocationKind, ParsedDocument, TableProbe
from app.ingestion.parsers import DocumentParseError

_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_WORD = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_DRAWING = "http://schemas.openxmlformats.org/drawingml/2006/main"
_PART_KIND: dict[str, LocationKind] = {"application/pdf": "page", _PPTX: "slide", _XLSX: "sheet"}


def _owned_elements(
    parent: ElementTree.Element, tag: str, *, boundaries: set[str]
) -> Iterator[ElementTree.Element]:
    """Traverse wrappers while keeping nested table content with its owner."""
    for child in parent:
        if child.tag == tag:
            yield child
        elif child.tag not in boundaries:
            yield from _owned_elements(child, tag, boundaries=boundaries)


def _ooxml_tables(data: bytes, *, presentation: bool) -> tuple[int, list[str]]:
    namespace = _DRAWING if presentation else _WORD
    # Native libraries discover parts through package relationships; legal
    # producers need not use the conventional document.xml / slideN.xml names.
    if presentation:
        from pptx import Presentation

        roots = [
            ElementTree.fromstring(slide.part.blob)
            for slide in Presentation(io.BytesIO(data)).slides
        ]
    else:
        from docx import Document

        roots = [ElementTree.fromstring(Document(io.BytesIO(data)).part.blob)]
    regions = 0
    cells: list[str] = []
    table_tag = f"{{{namespace}}}tbl"
    for root in roots:
        for table in root.iter(table_tag):
            regions += 1
            for row in _owned_elements(table, f"{{{namespace}}}tr", boundaries={table_tag}):
                for cell in _owned_elements(row, f"{{{namespace}}}tc", boundaries={table_tag}):
                    paragraphs = []
                    for paragraph in _owned_elements(
                        cell, f"{{{namespace}}}p", boundaries={table_tag}
                    ):
                        fragments = [
                            element.text or "" if element.tag == f"{{{namespace}}}t" else " "
                            for element in paragraph.iter()
                            if element.tag
                            in {
                                f"{{{namespace}}}t",
                                f"{{{namespace}}}br",
                                f"{{{namespace}}}cr",
                                f"{{{namespace}}}tab",
                            }
                        ]
                        paragraphs.append("".join(fragments))
                    text = " ".join(paragraphs)
                    if text.strip():
                        cells.append(text)
    return regions, cells


def _sheet_cells(data: bytes) -> tuple[int, list[str]]:
    from openpyxl import load_workbook

    formulas = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    try:
        cached = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            regions = 0
            cells: list[str] = []
            for sheet, values in zip(formulas.worksheets, cached.worksheets, strict=True):
                before = len(cells)
                for row, cached_row in zip(sheet.iter_rows(), values.iter_rows(), strict=True):
                    for cell, cached_cell in zip(row, cached_row, strict=True):
                        value = cached_cell.value if cell.data_type == "f" else cell.value
                        if value is None and cell.data_type == "f":
                            value = cell.value
                        if value is not None and str(value).strip():
                            cells.append(str(value))
                regions += len(cells) > before
            return regions, cells
        finally:
            cached.close()
    finally:
        formulas.close()


def build_extraction_diagnostics(
    data: bytes, *, mime_type: str, parsed: ParsedDocument
) -> ExtractionDiagnostics:
    normalized = mime_type.split(";", 1)[0].strip().lower()
    kind = _PART_KIND.get(normalized)
    blank = tuple(
        part.number
        for part in parsed.locations
        if not parsed.text[part.char_start : part.char_end].strip()
    )
    total = len(parsed.locations) if kind is not None else None
    replacements = parsed.text.count("\ufffd")
    controls = sum(
        unicodedata.category(character) == "Cc" and character not in "\t\n\r"
        for character in parsed.text
    )
    probe: TableProbe = "unavailable"
    regions: int | None = None
    count: int | None = None
    missing: int | None = None
    try:
        cells: list[str] = []
        if normalized in (_DOCX, _PPTX):
            probe = "native_tables"
            regions, cells = _ooxml_tables(data, presentation=normalized == _PPTX)
        elif normalized == _XLSX:
            probe = "sheet_cells"
            regions, cells = _sheet_cells(data)
        if probe != "unavailable":
            source = " ".join(parsed.text.split())
            count = len(cells)
            missing = sum(" ".join(cell.split()) not in source for cell in cells)
    except Exception as exc:  # noqa: BLE001 — fail closed on untrusted native bytes
        raise DocumentParseError(f"could not inspect native content: {type(exc).__name__}") from exc
    warnings: list[str] = []
    if blank:
        warnings.append("blank_native_parts")
    if replacements:
        warnings.append("replacement_characters")
    if controls:
        warnings.append("suspicious_controls")
    if missing:
        warnings.append("missing_native_table_text")
    if normalized == "application/pdf":
        warnings.append("pdf_table_coverage_unknown")
    return ExtractionDiagnostics(
        character_count=len(parsed.text),
        replacement_characters=replacements,
        suspicious_controls=controls,
        source_part_kind=kind,
        total_parts=total,
        parts_with_text=total - len(blank) if total is not None else None,
        blank_parts=blank,
        table_probe=probe,
        table_regions=regions,
        table_cells=count,
        missing_table_cells=missing,
        warnings=tuple(warnings),
    )
