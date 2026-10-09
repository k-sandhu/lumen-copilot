from __future__ import annotations

import io

import pytest
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Inches

from app.ingestion.chunking import chunk_text
from app.ingestion.parsers import DocumentParseError, parse_document

_PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def _serialize(presentation: Presentation) -> bytes:
    stream = io.BytesIO()
    presentation.save(stream)
    return stream.getvalue()


def _parse(presentation: Presentation) -> str:
    return parse_document(_serialize(presentation), mime_type=_PPTX)


def _add_textbox(shapes: object, text: str) -> object:
    textbox = shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    textbox.text_frame.text = text
    return textbox


def test_pptx_uses_source_title_and_numbers_untitled_content_slides() -> None:
    presentation = Presentation()
    titled = presentation.slides.add_slide(presentation.slide_layouts[0])
    titled.shapes.title.text = "Plan"
    _add_textbox(titled.shapes, "Visible plan detail")
    untitled = presentation.slides.add_slide(presentation.slide_layouts[6])
    _add_textbox(untitled.shapes, "Untitled slide detail")

    text = _parse(presentation)

    assert "Slide 1: Plan" in text
    assert "Visible plan detail" in text
    assert "Slide 2" in text
    assert "Untitled slide detail" in text
    assert text.index("Slide 1: Plan") < text.index("Slide 2")


def test_pptx_recursively_extracts_text_from_nested_group_shapes() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    _add_textbox(group.shapes, "Outer group text")
    nested_group = group.shapes.add_group_shape()
    _add_textbox(nested_group.shapes, "Nested group text")

    text = _parse(presentation)

    assert "Outer group text" in text
    assert "Nested group text" in text


def test_pptx_renders_table_rows_with_column_headers_and_units() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(5), Inches(2)).table
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Revenue (USD)"
    table.cell(1, 0).text = "West"
    table.cell(1, 1).text = "12"

    text = _parse(presentation)

    assert "[Table]" in text
    assert "Row 2: C1 [Region]=West | C2 [Revenue (USD)]=12" in text
    assert "[/Table]" in text


@pytest.mark.parametrize("cell_text", ["", " \t\n "], ids=["empty", "whitespace-only"])
def test_pptx_contentless_table_yields_no_text_or_chunks(cell_text: str) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(5), Inches(2)).table
    for row in table.rows:
        for cell in row.cells:
            cell.text = cell_text

    text = _parse(presentation)
    chunks = chunk_text(text, chunk_size=1000, overlap=100)

    assert text == ""
    assert chunks == []


def test_pptx_merged_table_text_stays_at_anchor_and_covered_cell_is_blank() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(5), Inches(2)).table
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Revenue (USD)"
    table.cell(1, 0).merge(table.cell(1, 1))
    table.cell(1, 0).text = "West"
    table.cell(1, 1).text = "Hidden covered text"
    reparsed = Presentation(io.BytesIO(_serialize(presentation)))
    covered = reparsed.slides[0].shapes[0].table.cell(1, 1)
    assert covered.is_spanned
    assert covered.text == "Hidden covered text"

    text = _parse(presentation)

    row = next(line for line in text.splitlines() if line.startswith("Row 2:"))
    assert "Hidden covered text" not in text
    assert row == "Row 2: C1 [Region]=West | C2 [Revenue (USD)]="
    assert row.count("West") == 1


def test_pptx_reads_notes_body_and_excludes_slide_number_placeholder() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    _add_textbox(slide.shapes, "Visible slide text")
    notes = slide.notes_slide
    notes.notes_text_frame.text = "Speaker-only detail"
    for placeholder in notes.placeholders:
        if placeholder.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
            placeholder.text = "777"

    text = _parse(presentation)

    assert "Notes (Slide 1):" in text
    assert "Speaker-only detail" in text
    assert "777" not in text


def test_pptx_does_not_create_notes_slide_when_notes_are_absent(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import pptx

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    _add_textbox(slide.shapes, "Visible text")
    assert not slide.has_notes_slide
    monkeypatch.setattr(pptx, "Presentation", lambda _data: presentation)

    text = parse_document(b"in-memory presentation", mime_type=_PPTX)

    assert "Visible text" in text
    assert not slide.has_notes_slide


def test_pptx_preserves_explicit_paragraph_line_break() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    paragraph = textbox.text_frame.paragraphs[0]
    paragraph.add_run().text = "First line"
    paragraph.add_line_break()
    paragraph.add_run().text = "Second line"

    text = _parse(presentation)

    assert "First line" in text
    assert "Second line" in text
    assert "\v" in text


@pytest.mark.parametrize("grouped", [False, True], ids=["ordinary", "grouped"])
@pytest.mark.parametrize(
    "source_text",
    [
        "First paragraph\nSecond paragraph",
        "First paragraph\n\nThird paragraph",
        "First paragraph\n\n\nFourth paragraph",
        "\n\nFirst paragraph\nLast paragraph\n\n",
    ],
    ids=["consecutive", "empty-middle", "multiple-empty-middle", "empty-edges"],
)
def test_pptx_preserves_complete_text_frame_paragraphs(source_text: str, grouped: bool) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shapes = slide.shapes.add_group_shape().shapes if grouped else slide.shapes
    _add_textbox(shapes, source_text)

    assert _parse(presentation) == f"Slide 1\n{source_text}"


@pytest.mark.parametrize("grouped", [False, True], ids=["ordinary", "grouped"])
def test_pptx_whitespace_only_paragraphs_yield_no_text_or_chunks(grouped: bool) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shapes = slide.shapes.add_group_shape().shapes if grouped else slide.shapes
    _add_textbox(shapes, "\n \t\n\n")

    text = _parse(presentation)

    assert text == ""
    assert chunk_text(text, chunk_size=1000, overlap=100) == []


def test_pptx_with_no_content_slides_returns_empty_text() -> None:
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])

    assert _parse(presentation) == ""


def test_corrupt_pptx_raises_typed_parse_error() -> None:
    with pytest.raises(DocumentParseError):
        parse_document(b"not a PPTX zip container", mime_type=_PPTX)
