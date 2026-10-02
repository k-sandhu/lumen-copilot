from __future__ import annotations

import io
import zipfile
from collections.abc import Callable
from xml.etree import ElementTree

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from app.ingestion.parsers import DocumentParseError, parse_document

_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _serialize(workbook: Workbook) -> bytes:
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _parse(workbook: Workbook) -> str:
    return parse_document(_serialize(workbook), mime_type=_XLSX)


def _rewrite_sheet(
    data: bytes,
    edit: Callable[[ElementTree.Element, str], None],
) -> bytes:
    source = io.BytesIO(data)
    output = io.BytesIO()
    with (
        zipfile.ZipFile(source) as archive,
        zipfile.ZipFile(output, mode="w", compression=zipfile.ZIP_DEFLATED) as updated,
    ):
        for item in archive.infolist():
            contents = archive.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                root = ElementTree.fromstring(contents)
                namespace = root.tag.partition("}")[0] + "}"
                edit(root, namespace)
                contents = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
            updated.writestr(item, contents)
    return output.getvalue()


def _set_cached_value(data: bytes, coordinate: str, value: str) -> bytes:
    def edit(root: ElementTree.Element, namespace: str) -> None:
        cell = root.find(f".//{namespace}c[@r='{coordinate}']")
        assert cell is not None
        cached = cell.find(f"{namespace}v")
        assert cached is not None
        cached.text = value

    return _rewrite_sheet(data, edit)


def test_xlsx_renders_each_nonempty_sheet_and_repeats_header_units() -> None:
    workbook = Workbook()
    sales = workbook.active
    sales.title = "Sales"
    sales.append(["Region", "Revenue (USD)"])
    sales.append(["West", 12])
    costs = workbook.create_sheet("Costs")
    costs.append(["Category", "Cost (CAD)"])
    costs.append(["Travel", 5])

    text = _parse(workbook)

    assert "Sheet: Sales" in text
    assert "Row 2 (Sheet Sales):" in text
    assert "A2 [Region]=West" in text
    assert "B2 [Revenue (USD)]=12" in text
    assert "Sheet: Costs" in text
    assert "Row 2 (Sheet Costs):" in text
    assert "A2 [Category]=Travel" in text
    assert "B2 [Cost (CAD)]=5" in text


def test_xlsx_keeps_column_coordinates_across_gaps_and_blank_cells() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Gaps"
    sheet["A1"] = "Left"
    sheet["C1"] = "Right"
    sheet["A2"] = "alpha"
    sheet["C2"] = "gamma"
    sheet["A4"] = "next row"
    sheet["C4"] = "still column C"

    text = _parse(workbook)

    row_two = next(line for line in text.splitlines() if line.startswith("Row 2 (Sheet Gaps):"))
    assert "A2 [Left]=alpha" in row_two
    assert "C2 [Right]=gamma" in row_two
    assert row_two.index("A2") < row_two.index("B2") < row_two.index("C2")
    assert "Row 3 (Sheet Gaps):" not in text
    row_four = next(line for line in text.splitlines() if line.startswith("Row 4 (Sheet Gaps):"))
    assert "C4" in row_four


def test_xlsx_preserves_zero_and_false_as_values() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Flags"
    sheet.append(["Count", "Enabled"])
    sheet.append([0, False])

    text = _parse(workbook)

    row = next(line for line in text.splitlines() if line.startswith("Row 2 (Sheet Flags):"))
    assert "A2 [Count]=0" in row
    assert "B2 [Enabled]=False" in row


def test_xlsx_retains_non_general_percent_number_format() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Rates"
    sheet.append(["Completion"])
    sheet["A2"] = 0.25
    sheet["A2"].number_format = "0.00%"

    text = _parse(workbook)

    assert "A2 [Completion]=0.25 [format=0.00%]" in text


def test_xlsx_reports_merged_ranges_and_leaves_covered_cells_blank() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Merged"
    sheet.merge_cells("A1:B1")
    sheet["A1"] = "Combined heading"
    sheet["A2"] = "Detail"
    sheet["B2"] = "Value"
    sheet["A3"] = "item"
    sheet["B3"] = 7

    text = _parse(workbook)

    assert "Merged cells: A1:B1" in text
    assert "A1=Combined heading" in text
    assert "B1=Combined heading" not in text
    assert "B1=" in text  # the covered position is retained, without a second value


@pytest.mark.parametrize("formula", [False, True], ids=["numeric", "formula"])
@pytest.mark.parametrize("anchor", [None, "Heading"], ids=["empty-anchor", "heading-anchor"])
def test_xlsx_ignores_stored_merged_covered_values_before_selecting_headers(
    formula: bool,
    anchor: str | None,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Covered"
    sheet.merge_cells("A1:B1")
    sheet["A1"] = anchor
    sheet.append(["Item", 5])
    sheet.append(["Next", 7])

    def edit(root: ElementTree.Element, namespace: str) -> None:
        row = root.find(f".//{namespace}row[@r='1']")
        assert row is not None
        cell = ElementTree.SubElement(row, f"{namespace}c", r="B1", t="n")
        if formula:
            ElementTree.SubElement(cell, f"{namespace}f").text = "SUM(B2:B3)"
        ElementTree.SubElement(cell, f"{namespace}v").text = "99"

    data = _rewrite_sheet(_serialize(workbook), edit)
    expected = (
        "Sheet: Covered\nMerged cells: A1:B1\n"
        "Row 1 (Sheet Covered): A1=Heading | B1=\n"
        "Row 2 (Sheet Covered): A2 [Heading]=Item | B2=5\n"
        "Row 3 (Sheet Covered): A3 [Heading]=Next | B3=7"
        if anchor is not None
        else "Sheet: Covered\nMerged cells: A1:B1\n"
        "Row 2 (Sheet Covered): A2=Item | B2=5\n"
        "Row 3 (Sheet Covered): A3 [Item]=Next | B3 [5]=7"
    )

    assert parse_document(data, mime_type=_XLSX) == expected


def test_xlsx_skips_rows_containing_only_stored_merged_covered_values() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Covered"
    sheet.merge_cells("A1:B2")
    sheet["A1"] = "Heading"
    sheet["A3"] = "Item"
    sheet["B3"] = 5

    def edit(root: ElementTree.Element, namespace: str) -> None:
        row = root.find(f".//{namespace}row[@r='2']")
        assert row is not None
        cell = ElementTree.SubElement(row, f"{namespace}c", r="A2", t="n")
        ElementTree.SubElement(cell, f"{namespace}f").text = "SUM(B3)"
        # Even an uncached covered formula cannot make this row nonblank.
        ElementTree.SubElement(cell, f"{namespace}v")
        cell = ElementTree.SubElement(row, f"{namespace}c", r="B2", t="n")
        ElementTree.SubElement(cell, f"{namespace}v").text = "99"

    data = _rewrite_sheet(_serialize(workbook), edit)

    assert parse_document(data, mime_type=_XLSX) == (
        "Sheet: Covered\nMerged cells: A1:B2\n"
        "Row 1 (Sheet Covered): A1=Heading | B1=\n"
        "Row 3 (Sheet Covered): A3 [Heading]=Item | B3=5"
    )


@pytest.mark.parametrize("reference", ["B2:A1", "A:B", "A0:B1", "not-a-range"])
def test_xlsx_rejects_invalid_merged_ranges(reference: str) -> None:
    workbook = Workbook()
    workbook.active["A1"] = "Heading"
    workbook.active.merge_cells("A1:B1")

    def edit(root: ElementTree.Element, namespace: str) -> None:
        merge = root.find(f".//{namespace}mergeCell")
        assert merge is not None
        merge.set("ref", reference)

    with pytest.raises(DocumentParseError):
        parse_document(_rewrite_sheet(_serialize(workbook), edit), mime_type=_XLSX)


def test_xlsx_marks_formula_when_cached_value_is_unavailable() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Totals"
    sheet.append(["Amount"])
    sheet.append([2])
    sheet.append([3])
    sheet["A4"] = "=SUM(A2:A3)"

    text = _parse(workbook)

    row = next(line for line in text.splitlines() if line.startswith("Row 4 (Sheet Totals):"))
    assert "A4 [Amount]" in row
    assert "[formula==SUM(A2:A3); cached value unavailable]" in row


def test_xlsx_marks_formula_and_supplied_cached_value_as_freshness_unknown() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Totals"
    sheet.append(["Amount"])
    sheet.append([2])
    sheet.append([3])
    sheet["A4"] = "=SUM(A2:A3)"
    data = _set_cached_value(_serialize(workbook), "A4", "5")

    text = parse_document(data, mime_type=_XLSX)

    row = next(line for line in text.splitlines() if line.startswith("Row 4 (Sheet Totals):"))
    assert "A4 [Amount]=5" in row
    assert "[formula==SUM(A2:A3); cached value supplied; freshness unknown]" in row


@pytest.mark.parametrize("cached_value", [None, "5"], ids=["uncached", "cached"])
def test_xlsx_renders_array_formula_expression_and_range_deterministically(
    cached_value: str | None,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Arrays"
    sheet["A1"] = "Total"
    sheet["A2"] = ArrayFormula(ref="A2:A3", text="=B2:B3+C2:C3")
    data = _serialize(workbook)
    if cached_value is not None:
        data = _set_cached_value(data, "A2", cached_value)
    cache = (
        "cached value unavailable"
        if cached_value is None
        else "cached value supplied; freshness unknown"
    )
    expected = (
        "Sheet: Arrays\nRow 1 (Sheet Arrays): A1=Total\n"
        f"Row 2 (Sheet Arrays): A2 [Total]={cached_value or ''} "
        f"[formula==B2:B3+C2:C3; array range=A2:A3; {cache}]"
    )

    assert parse_document(data, mime_type=_XLSX) == expected
    assert parse_document(data, mime_type=_XLSX) == expected


@pytest.mark.parametrize("cached_value", [None, "5"], ids=["uncached", "cached"])
@pytest.mark.parametrize("two_inputs", [False, True], ids=["one-input", "two-inputs"])
def test_xlsx_renders_data_table_formula_attributes_deterministically(
    cached_value: str | None,
    two_inputs: bool,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tables"
    sheet["A1"] = "Result"
    sheet["A2"] = (
        DataTableFormula(
            ref="A2:A3", dt2D=True, dtr=True, r1="B1", r2="C1", del1=True, del2=True, ca=True
        )
        if two_inputs
        else DataTableFormula(ref="A2:A3", r1="B1")
    )
    data = _serialize(workbook)
    if cached_value is not None:
        data = _set_cached_value(data, "A2", cached_value)
    attributes = (
        "ref=A2:A3; dt2D=1; dtr=1; r1=B1; r2=C1; del1=1; del2=1; ca=1"
        if two_inputs
        else "ref=A2:A3; r1=B1"
    )
    cache = (
        "cached value unavailable"
        if cached_value is None
        else "cached value supplied; freshness unknown"
    )
    expected = (
        "Sheet: Tables\nRow 1 (Sheet Tables): A1=Result\n"
        f"Row 2 (Sheet Tables): A2 [Result]={cached_value or ''} "
        f"[formula=dataTable; {attributes}; {cache}]"
    )

    assert parse_document(data, mime_type=_XLSX) == expected
    assert parse_document(data, mime_type=_XLSX) == expected


@pytest.mark.parametrize("formula_kind", ["ordinary", "array", "data-table"])
@pytest.mark.parametrize(
    ("cell_type", "cache_text", "supplied", "rendered_value"),
    [
        ("str", "", True, ""),
        ("str", None, False, ""),
        ("n", "", False, ""),
        (None, "", False, ""),
        (None, None, False, ""),
        ("n", "0", True, "0"),
        ("b", "0", True, "False"),
        ("e", "#DIV/0!", True, "#DIV/0!"),
        ("str", "ready", True, "ready"),
    ],
    ids=[
        "empty-string",
        "absent-string",
        "empty-numeric",
        "empty-default",
        "absent-default",
        "zero",
        "false",
        "error",
        "string",
    ],
)
def test_xlsx_formula_cache_availability_preserves_source_presence_and_type(
    formula_kind: str,
    cell_type: str | None,
    cache_text: str | None,
    supplied: bool,
    rendered_value: str,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Caches"
    sheet["A1"] = "Status"
    if formula_kind == "ordinary":
        sheet["A2"] = '=IF(FALSE,1,"")'
        expression = '=IF(FALSE,1,"")'
    elif formula_kind == "array":
        sheet["A2"] = ArrayFormula(ref="A2:A3", text='=IF(FALSE,1,"")')
        expression = '=IF(FALSE,1,""); array range=A2:A3'
    else:
        sheet["A2"] = DataTableFormula(ref="A2:A3", r1="B1")
        expression = "dataTable; ref=A2:A3; r1=B1"

    def edit(root: ElementTree.Element, namespace: str) -> None:
        cell = root.find(f".//{namespace}c[@r='A2']")
        assert cell is not None
        if cell_type is None:
            cell.attrib.pop("t", None)
        else:
            cell.set("t", cell_type)
        cached = cell.find(f"{namespace}v")
        assert cached is not None
        if cache_text is None:
            cell.remove(cached)
        else:
            cached.text = cache_text

    data = _rewrite_sheet(_serialize(workbook), edit)
    cache = "cached value supplied; freshness unknown" if supplied else "cached value unavailable"
    expected = (
        "Sheet: Caches\nRow 1 (Sheet Caches): A1=Status\n"
        f"Row 2 (Sheet Caches): A2 [Status]={rendered_value} [formula={expression}; {cache}]"
    )

    assert parse_document(data, mime_type=_XLSX) == expected
    assert parse_document(data, mime_type=_XLSX) == expected


def test_xlsx_with_no_nonblank_cells_returns_empty_text() -> None:
    workbook = Workbook()
    workbook.active.title = "Empty"

    assert _parse(workbook) == ""


def test_corrupt_xlsx_raises_typed_parse_error() -> None:
    with pytest.raises(DocumentParseError):
        parse_document(b"not an XLSX workbook", mime_type=_XLSX)
