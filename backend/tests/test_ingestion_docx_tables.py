from __future__ import annotations

import io

import docx
import pytest
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import _Cell

import app.ingestion.parsers as parser_module
from app.ingestion.chunking import chunk_text
from app.ingestion.parsers import DocumentParseError, parse_document

_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _serialize(document: docx.document.Document) -> bytes:
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _parse(document: docx.document.Document) -> str:
    return parse_document(_serialize(document), mime_type=_DOCX)


def _make_table_document(
    rows: int, columns: int
) -> tuple[docx.document.Document, docx.table.Table]:
    document = docx.Document()
    table = document.add_table(rows=rows, cols=columns)
    return document, table


def test_docx_extracts_table_only_facts_with_header_labels_and_units() -> None:
    document, table = _make_table_document(2, 2)
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Revenue (USD)"
    table.cell(1, 0).text = "West"
    table.cell(1, 1).text = "12"

    text = _parse(document)

    assert "[Table]" in text
    assert "Row 1:" in text and "C1=Region" in text
    assert "C2=Revenue (USD)" in text
    assert "Row 2:" in text
    assert "C1 [Region]=West" in text
    assert "C2 [Revenue (USD)]=12" in text
    assert "[/Table]" in text


def test_docx_preserves_paragraph_table_adjacency_order_and_paragraph_text() -> None:
    document = docx.Document()
    document.add_paragraph("Before table, unchanged.")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Middle table fact"
    document.add_paragraph("After table, unchanged.")

    text = _parse(document)

    before = text.index("Before table, unchanged.")
    table_start = text.index("[Table]")
    fact = text.index("Middle table fact")
    table_end = text.index("[/Table]")
    after = text.index("After table, unchanged.")
    assert before < table_start < fact < table_end < after
    assert "Before table, unchanged." in text
    assert "After table, unchanged." in text


def test_docx_table_keeps_blank_middle_grid_column() -> None:
    document, table = _make_table_document(2, 3)
    table.cell(0, 0).text = "Left"
    table.cell(0, 1).text = ""
    table.cell(0, 2).text = "Right"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = ""
    table.cell(1, 2).text = "B"

    text = _parse(document)

    row = next(line for line in text.splitlines() if line.startswith("Row 2:"))
    assert "C1 [Left]=A" in row
    assert "C2" in row
    assert "C3 [Right]=B" in row


def test_docx_horizontally_merged_cells_keep_their_grid_positions() -> None:
    document, table = _make_table_document(2, 3)
    table.cell(0, 0).text = "Department"
    table.cell(0, 1).text = "Quarter"
    table.cell(0, 2).text = "Amount"
    merged = table.cell(1, 0).merge(table.cell(1, 1))
    merged.text = "North"
    table.cell(1, 2).text = "8"

    text = _parse(document)

    row = next(line for line in text.splitlines() if line.startswith("Row 2:"))
    assert "C1 [Department]=North" in row
    assert "C2 [Quarter]=North" in row
    assert "C3 [Amount]=8" in row


def test_docx_vertically_merged_cells_repeat_value_at_occupied_positions() -> None:
    document, table = _make_table_document(3, 2)
    table.cell(0, 0).text = "Team"
    table.cell(0, 1).text = "Owner"
    merged = table.cell(1, 0).merge(table.cell(2, 0))
    merged.text = "Platform"
    table.cell(1, 1).text = "Ari"
    table.cell(2, 1).text = "Bo"

    text = _parse(document)

    row_two = next(line for line in text.splitlines() if line.startswith("Row 2:"))
    row_three = next(line for line in text.splitlines() if line.startswith("Row 3:"))
    assert "C1 [Team]=Platform" in row_two
    assert "C1 [Team]=Platform" in row_three
    assert "C2 [Owner]=Ari" in row_two
    assert "C2 [Owner]=Bo" in row_three


def test_docx_horizontal_merge_marks_repeated_numeric_cell_origin() -> None:
    document, table = _make_table_document(2, 3)
    table.cell(0, 0).text = "Quantity"
    table.cell(0, 1).text = "Unit price"
    table.cell(0, 2).text = "Total"
    merged = table.cell(1, 0).merge(table.cell(1, 1))
    merged.text = "120"
    table.cell(1, 2).text = "240"

    text = _parse(document)

    row = next(line for line in text.splitlines() if line.startswith("Row 2:"))
    assert "C1 [Quantity]=120" in row
    assert "C2 [Unit price]=120 [merged from R2C1]" in row
    assert "C1 [Quantity]=120 [merged from" not in row
    assert "C3 [Total]=240 [merged from" not in row


def test_docx_vertical_merge_marks_repeated_numeric_cell_origin() -> None:
    document, table = _make_table_document(3, 2)
    table.cell(0, 0).text = "Amount"
    table.cell(0, 1).text = "Period"
    merged = table.cell(1, 0).merge(table.cell(2, 0))
    merged.text = "75"
    table.cell(1, 1).text = "Q1"
    table.cell(2, 1).text = "Q2"

    text = _parse(document)

    row_two = next(line for line in text.splitlines() if line.startswith("Row 2:"))
    row_three = next(line for line in text.splitlines() if line.startswith("Row 3:"))
    assert "C1 [Amount]=75" in row_two
    assert "C1 [Amount]=75 [merged from R2C1]" in row_three
    assert "C2 [Period]=Q2 [merged from" not in row_three


def test_docx_equal_unmerged_numeric_cells_are_not_marked_as_merged() -> None:
    document, table = _make_table_document(2, 2)
    table.cell(0, 0).text = "Count A"
    table.cell(0, 1).text = "Count B"
    table.cell(1, 0).text = "5"
    table.cell(1, 1).text = "5"

    text = _parse(document)

    row = next(line for line in text.splitlines() if line.startswith("Row 2:"))
    assert "C1 [Count A]=5" in row
    assert "C2 [Count B]=5" in row
    assert "[merged from" not in row


def test_docx_recursively_extracts_nested_table_text() -> None:
    document, table = _make_table_document(1, 1)
    cell = table.cell(0, 0)
    cell.paragraphs[0].text = "Outer cell"
    nested = cell.add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested secret code"

    text = _parse(document)

    assert "Outer cell" in text
    assert "Nested secret code" in text


def test_docx_preserves_paragraph_only_text_exactly() -> None:
    document = docx.Document()
    document.add_paragraph("First paragraph.")
    document.add_paragraph("Second paragraph.")

    assert _parse(document) == "First paragraph.\nSecond paragraph."


def test_corrupt_docx_still_raises_typed_parse_error() -> None:
    with pytest.raises(DocumentParseError):
        parse_document(b"not a DOCX zip container", mime_type=_DOCX)


def test_docx_long_vertical_merge_retains_last_value_and_origin() -> None:
    """R1-001: resolve 1,000 continuations without recursive row.cells."""
    document, table = _make_table_document(1000, 2)
    for number, row in enumerate(table.rows):
        group = _Cell(row._tr.tc_lst[0], table)
        group.text = "Group" if number == 0 else ""
        merge = OxmlElement("w:vMerge")
        merge.set(qn("w:val"), "restart" if number == 0 else "continue")
        group._tc.get_or_add_tcPr().append(merge)
        _Cell(row._tr.tc_lst[1], table).text = f"unique value {number}"

    text = _parse(document)

    last_row = next(line for line in text.splitlines() if line.startswith("Row 1000:"))
    assert "C1 [Group]=Group [merged from R1C1]" in last_row
    assert "C2 [unique value 0]=unique value 999" in last_row
    assert text.count("unique value 999") == 1


def test_docx_vertical_span_with_omitted_columns_keeps_absolute_grid_origins() -> None:
    """R1-001: iterative origins must include omissions and horizontal spans."""
    document, table = _make_table_document(3, 4)
    for column, name in enumerate(("Before", "Group", "Alias", "After")):
        table.cell(0, column).text = name
    table.cell(1, 1).merge(table.cell(2, 2)).text = "Shared"
    for row in table.rows[1:]:
        row._tr.remove(row._tr.tc_lst[-1])
        row._tr.remove(row._tr.tc_lst[0])
        for name in ("gridBefore", "gridAfter"):
            omission = OxmlElement(f"w:{name}")
            omission.set(qn("w:val"), "1")
            row._tr.get_or_add_trPr().append(omission)

    text = _parse(document)

    last_row = next(line for line in text.splitlines() if line.startswith("Row 3:"))
    assert last_row == (
        "Row 3: C1 [Before]= | C2 [Group]=Shared [merged from R2C2] | "
        "C3 [Alias]=Shared [merged from R2C2] | C4 [After]="
    )


def _nested_merged_document(levels: int) -> docx.document.Document:
    document, table = _make_table_document(1, 2)
    cell = table.cell(0, 0).merge(table.cell(0, 1))
    for _ in range(levels - 1):
        table = cell.add_table(rows=1, cols=2)
        cell = table.cell(0, 0).merge(table.cell(0, 1))
    cell.text = "unique leaf fact"
    return document


def test_docx_nested_merged_aliases_do_not_multiply_leaf_facts() -> None:
    """R1-002: 17 merged levels used to produce 131,072 leaf copies."""
    text = _parse(_nested_merged_document(17))

    # The deepest scalar still repeats once for its horizontal alias; every
    # ancestor references its nested origin rather than copying the subtree.
    assert text.count("unique leaf fact") == 2
    assert text.count("[Table]") == 17
    assert text.count("[nested tables at R1C1]") == 16
    assert text.count("[merged from R1C1]") == 17
    assert len(chunk_text(text, chunk_size=1000, overlap=150)) <= 3


def test_docx_nested_first_row_headers_and_aliases_keep_scalar_context_only() -> None:
    document, table = _make_table_document(2, 2)
    origin = table.cell(0, 0).merge(table.cell(0, 1))
    origin.paragraphs[0].text = "Portfolio (USD)"
    origin.add_table(rows=1, cols=1).cell(0, 0).text = "unique nested amount"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "B"

    text = _parse(document)

    assert text.count("unique nested amount") == 1
    assert "[nested tables at R1C1] [merged from R1C1]" in text
    assert "C1 [Portfolio (USD) [nested tables at R1C1]]=A" in text
    assert "C2 [Portfolio (USD) [nested tables at R1C1]]=B" in text


def test_docx_output_budget_counts_repeated_headers_and_merge_values(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(parser_module, "_DOCX_MAX_OUTPUT_CHARS", 200, raising=False)
    document, table = _make_table_document(3, 2)
    table.cell(0, 0).text = "Long source header (USD)"
    table.cell(0, 1).text = "Second source header (USD)"
    table.cell(1, 0).merge(table.cell(2, 1)).text = "120"

    with pytest.raises(DocumentParseError, match="DOCX extraction output limit"):
        _parse(document)


@pytest.mark.parametrize("expansion", ["span", "before", "after"])
def test_docx_work_budget_precedes_grid_expansion(
    monkeypatch: pytest.MonkeyPatch, expansion: str
) -> None:
    monkeypatch.setattr(parser_module, "_DOCX_MAX_WORK_UNITS", 20, raising=False)
    document, table = _make_table_document(1, 1)
    if expansion == "span":
        table.cell(0, 0)._tc.grid_span = 21
    else:
        omission = OxmlElement("w:gridBefore" if expansion == "before" else "w:gridAfter")
        omission.set(qn("w:val"), "21")
        table.rows[0]._tr.get_or_add_trPr().append(omission)

    def guarded_range(start: int, stop: int) -> range:
        assert stop - start <= 20, "grid expanded before checking the work budget"
        return range(start, stop)

    monkeypatch.setattr(parser_module, "range", guarded_range, raising=False)
    with pytest.raises(DocumentParseError, match="DOCX extraction work limit"):
        _parse(document)


def test_docx_nested_depth_budget_fails_with_typed_error(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(parser_module, "_DOCX_MAX_TABLE_DEPTH", 4, raising=False)

    with pytest.raises(DocumentParseError, match="DOCX extraction depth limit"):
        _parse(_nested_merged_document(5))


def test_docx_scalar_budget_stops_before_reading_more_paragraphs(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """R1-002: stop at the boundary, before collecting an oversized scalar."""
    monkeypatch.setattr(parser_module, "_DOCX_MAX_OUTPUT_CHARS", 80, raising=False)
    document, table = _make_table_document(1, 1)
    cell = table.cell(0, 0)
    cell.text = "a" * 50
    cell.add_paragraph("b" * 50)
    cell.add_paragraph("unvisited sentinel")
    text_property = docx.text.paragraph.Paragraph.text
    observed: list[str] = []

    def read_text(paragraph: docx.text.paragraph.Paragraph) -> str:
        text = text_property.fget(paragraph)
        assert text != "unvisited sentinel", "continued reading after the scalar limit"
        observed.append(text)
        return text

    monkeypatch.setattr(docx.text.paragraph.Paragraph, "text", property(read_text))
    with pytest.raises(DocumentParseError, match="DOCX extraction output limit"):
        _parse(document)
    assert observed == ["a" * 50, "b" * 50]
