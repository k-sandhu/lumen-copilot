from __future__ import annotations

import io
import uuid
import xml.etree.ElementTree as ET
from collections.abc import AsyncIterator, Sequence
from importlib import import_module
from zipfile import ZipFile

import pytest
import pytest_asyncio

_PDF = "application/pdf"
_PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _make_pdf_pages_with_blank_middle() -> bytes:
    from pypdf import PdfReader, PdfWriter

    from tests.test_ingestion_parsers import _make_pdf

    writer = PdfWriter()
    writer.append(PdfReader(io.BytesIO(_make_pdf("First page sentinel"))))
    writer.add_blank_page(width=612, height=792)
    writer.append(PdfReader(io.BytesIO(_make_pdf("Final page sentinel"))))
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def _make_presentation_with_blank_second_slide() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[0])
    first.shapes.title.text = "Project plan"
    body = first.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1))
    body.text_frame.text = "Milestone detail"
    presentation.slides.add_slide(presentation.slide_layouts[6])
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def _make_workbook_with_blank_middle_sheet() -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    north = workbook.active
    north.title = "North"
    north.append(["Region", "Revenue"])
    north.append(["North", 12])
    workbook.create_sheet("Blank Middle")
    south = workbook.create_sheet("South")
    south.append(["Region", "Revenue"])
    south.append(["South", 18])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _make_workbook_with_empty_shared_string_sheet(position: int) -> bytes:
    """Author a real empty shared string; openpyxl's writer omits empty cells."""
    from openpyxl import Workbook

    workbook = Workbook()
    values = ["alpha", "omega"]
    values.insert(position, "")
    for number, value in enumerate(values, start=1):
        sheet = workbook.active if number == 1 else workbook.create_sheet()
        sheet.title = f"Sheet {number}"
        if value:
            sheet["A1"] = value
    original = io.BytesIO()
    workbook.save(original)
    workbook.close()

    spreadsheet_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    output = io.BytesIO()
    with ZipFile(original) as source, ZipFile(output, "w") as target:
        for entry in source.infolist():
            payload = source.read(entry.filename)
            if entry.filename == f"xl/worksheets/sheet{position + 1}.xml":
                payload = (
                    f'<worksheet xmlns="{spreadsheet_ns}"><sheetData><row r="1">'
                    '<c r="A1" t="s"><v>0</v></c></row></sheetData></worksheet>'
                ).encode()
            elif entry.filename == "[Content_Types].xml":
                root = ET.fromstring(payload)
                ET.SubElement(
                    root,
                    "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
                    PartName="/xl/sharedStrings.xml",
                    ContentType="application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sharedStrings+xml",
                )
                payload = ET.tostring(root)
            elif entry.filename == "xl/_rels/workbook.xml.rels":
                root = ET.fromstring(payload)
                ET.SubElement(
                    root,
                    "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                    Id="emptySharedString",
                    Type="http://schemas.openxmlformats.org/officeDocument/2006/"
                    "relationships/sharedStrings",
                    Target="sharedStrings.xml",
                )
                payload = ET.tostring(root)
            target.writestr(entry, payload)
        target.writestr(
            "xl/sharedStrings.xml",
            f'<sst xmlns="{spreadsheet_ns}" count="1" uniqueCount="1"><si><t/></si></sst>',
        )
    return output.getvalue()


def _assert_location_bounds(parsed: object) -> None:
    text = parsed.text
    for location in parsed.locations:
        assert 0 <= location.char_start <= location.char_end <= len(text)


@pytest.mark.parametrize(
    ("mime_type", "data"),
    [(_PDF, b"not a PDF"), (_PPTX, b"not a PPTX"), (_XLSX, b"not an XLSX")],
)
def test_corrupt_native_formats_keep_typed_parse_errors(mime_type: str, data: bytes) -> None:
    from app.ingestion.parsers import DocumentParseError, parse_document_with_locations

    with pytest.raises(DocumentParseError):
        parse_document_with_locations(data, mime_type=mime_type)


def test_pdf_locations_cover_named_pages_and_preserve_blank_page_span() -> None:
    from app.ingestion.parsers import parse_document, parse_document_with_locations

    data = _make_pdf_pages_with_blank_middle()
    parsed = parse_document_with_locations(data, mime_type=_PDF)

    assert parse_document(data, mime_type=_PDF) == parsed.text
    assert [item.kind for item in parsed.locations] == ["page", "page", "page"]
    assert [item.number for item in parsed.locations] == [1, 2, 3]
    assert "Firstpagesentinel" in parsed.text.replace(" ", "")
    assert "Finalpagesentinel" in parsed.text.replace(" ", "")
    assert parsed.locations[0].char_start < parsed.locations[0].char_end
    assert parsed.locations[1].char_start == parsed.locations[1].char_end
    assert parsed.locations[2].char_start < parsed.locations[2].char_end
    _assert_location_bounds(parsed)


def test_pptx_locations_cover_slides_and_preserve_blank_slide_span() -> None:
    from app.ingestion.parsers import parse_document, parse_document_with_locations

    data = _make_presentation_with_blank_second_slide()
    parsed = parse_document_with_locations(data, mime_type=_PPTX)

    assert parse_document(data, mime_type=_PPTX) == parsed.text
    assert [item.kind for item in parsed.locations] == ["slide", "slide"]
    assert [item.number for item in parsed.locations] == [1, 2]
    assert any("Project plan" in item.name for item in parsed.locations)
    assert "Milestone detail" in parsed.text
    assert parsed.locations[1].char_start == parsed.locations[1].char_end
    _assert_location_bounds(parsed)


def test_xlsx_locations_use_sheet_names_and_preserve_blank_middle_sheet_span() -> None:
    from app.ingestion.parsers import parse_document, parse_document_with_locations

    data = _make_workbook_with_blank_middle_sheet()
    parsed = parse_document_with_locations(data, mime_type=_XLSX)

    assert parse_document(data, mime_type=_XLSX) == parsed.text
    assert [item.kind for item in parsed.locations] == ["sheet", "sheet", "sheet"]
    assert [item.name for item in parsed.locations] == ["North", "Blank Middle", "South"]
    assert [item.number for item in parsed.locations] == [1, 2, 3]
    assert "North" in parsed.text and "South" in parsed.text
    assert parsed.locations[1].char_start == parsed.locations[1].char_end
    _assert_location_bounds(parsed)


@pytest.mark.parametrize("with_locations", [False, True], ids=["legacy", "mapped"])
@pytest.mark.parametrize(
    ("position", "expected_text", "spans"),
    [
        (0, "\n\nalpha\n\nomega", [(0, 0), (2, 7), (9, 14)]),
        (1, "alpha\n\n\n\nomega", [(0, 5), (7, 7), (9, 14)]),
        (2, "alpha\n\nomega\n\n", [(0, 5), (7, 12), (14, 14)]),
    ],
    ids=["leading", "middle", "trailing"],
)
def test_xlsx_empty_shared_string_sheet_preserves_golden_rendering_and_offsets(
    position: int, expected_text: str, spans: list[tuple[int, int]], with_locations: bool
) -> None:
    from openpyxl import load_workbook

    from app.ingestion.chunking import chunk_text
    from app.ingestion.parsers import parse_document, parse_document_with_locations

    data = _make_workbook_with_empty_shared_string_sheet(position)
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        assert workbook.worksheets[position]["A1"].value == ""
    finally:
        workbook.close()

    if not with_locations:
        assert parse_document(data, mime_type=_XLSX) == expected_text
        return

    parsed = parse_document_with_locations(data, mime_type=_XLSX)
    assert parsed.text == expected_text
    assert [item.kind for item in parsed.locations] == ["sheet"] * 3
    assert [item.number for item in parsed.locations] == [1, 2, 3]
    assert [item.name for item in parsed.locations] == ["Sheet 1", "Sheet 2", "Sheet 3"]
    assert [(item.char_start, item.char_end) for item in parsed.locations] == spans
    values = ["alpha", "omega"]
    values.insert(position, "")
    assert [parsed.text[start:end] for start, end in spans] == values

    chunks = chunk_text(parsed.text, chunk_size=256, overlap=0)
    assert len(chunks) == 1
    chunk = chunks[0]
    assert parsed.text[chunk.char_start : chunk.char_end] == chunk.text
    assert [
        item.number for item in parsed.locations_for_span(chunk.char_start, chunk.char_end)
    ] == [number for number in range(1, 4) if number != position + 1]


def test_chunk_crossing_pdf_pages_has_exact_text_and_both_locations() -> None:
    from app.ingestion.chunking import chunk_text
    from app.ingestion.parsers import parse_document_with_locations

    parsed = parse_document_with_locations(_make_pdf_pages_with_blank_middle(), mime_type=_PDF)
    chunks = chunk_text(parsed.text, chunk_size=256, overlap=0)

    assert len(chunks) == 1
    chunk = chunks[0]
    assert parsed.text[chunk.char_start : chunk.char_end] == chunk.text
    mapped = parsed.locations_for_span(chunk.char_start, chunk.char_end)
    assert [item.number for item in mapped] == [1, 3]
    assert all(item.char_start < item.char_end for item in mapped)


@pytest.mark.parametrize(
    "overrides",
    [
        {"number": 0},
        {"char_end": -1},
        {"char_start": 2, "char_end": 1},
    ],
)
def test_source_location_rejects_invalid_numbers_and_spans(
    overrides: dict[str, int],
) -> None:
    source_location = import_module("app.domain.ingestion").SourceLocation
    values: dict[str, object] = {
        "kind": "page",
        "name": "Page 1",
        "number": 1,
        "char_start": 0,
        "char_end": 1,
        **overrides,
    }

    with pytest.raises(ValueError):
        source_location(**values)


def test_locations_for_span_rejects_negative_offsets() -> None:
    parse_document_with_locations = import_module(
        "app.ingestion.parsers"
    ).parse_document_with_locations
    parsed = parse_document_with_locations(_make_pdf_pages_with_blank_middle(), mime_type=_PDF)

    with pytest.raises(ValueError):
        parsed.locations_for_span(-1, 1)


@pytest_asyncio.fixture
async def sqlite_engine() -> AsyncIterator[None]:
    from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
    from sqlalchemy.pool import StaticPool

    import app.db.models  # noqa: F401 — register tables with Base.metadata
    import app.db.session as db_session
    from app.db.base import Base

    engine = create_async_engine(
        "sqlite+aiosqlite://",
        poolclass=StaticPool,
        connect_args={"check_same_thread": False},
    )
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    previous_engine = db_session._engine
    previous_sessionmaker = db_session._sessionmaker
    db_session._engine = engine
    db_session._sessionmaker = async_sessionmaker(bind=engine, expire_on_commit=False)
    try:
        yield
    finally:
        db_session._engine = previous_engine
        db_session._sessionmaker = previous_sessionmaker
        await engine.dispose()


class _NoopIndexStore:
    @classmethod
    def from_settings(cls, _settings: object) -> _NoopIndexStore:
        return cls()

    async def ensure_index(self) -> None:
        return None

    async def upsert_chunks(self, _chunks: Sequence[object], *, refresh: bool = False) -> None:
        return None

    async def delete_document(
        self, *, tenant_id: uuid.UUID, document_id: uuid.UUID, refresh: bool = False
    ) -> None:
        return None

    async def aclose(self) -> None:
        return None


async def test_ingestion_persists_extracted_text_and_matching_part_maps(
    sqlite_engine: None, monkeypatch: pytest.MonkeyPatch
) -> None:
    parse_document_with_locations = import_module(
        "app.ingestion.parsers"
    ).parse_document_with_locations
    db_session = import_module("app.db.session")
    index_sync = import_module("app.tasks.index_sync")
    repositories = import_module("app.db.repositories")
    ingest = import_module("app.tasks.ingest")
    from tests.test_ingestion_task import (
        _FakeGateway,
        _FakeObjectStore,
        _seed_document,
        _settings,
    )

    monkeypatch.setattr(index_sync, "OpenSearchStore", _NoopIndexStore)
    data = _make_pdf_pages_with_blank_middle()
    parsed = parse_document_with_locations(data, mime_type=_PDF)
    tenant_id, document_id = await _seed_document(mime_type=_PDF, key="pending")
    store = _FakeObjectStore()
    async with db_session.session_scope() as session:
        document = await repositories.DocumentRepository(session, tenant_id).get(document_id)
        assert document is not None
        store.put(str(tenant_id), document.storage_key, data)

    await ingest.ingest_document_async(
        tenant_id,
        document_id,
        settings=_settings(),
        object_store=store,
        gateway=_FakeGateway(),
    )

    async with db_session.session_scope() as session:
        document = await repositories.DocumentRepository(session, tenant_id).get(document_id)
        assert document is not None
        assert document.source_text == parsed.text
        assert document.source_locations == parsed.locations
        chunks = await repositories.ChunkRepository(session, tenant_id).list_for_document(
            document_id
        )
        assert chunks
        for chunk in chunks:
            assert chunk.source_locations == parsed.locations_for_span(
                chunk.char_start, chunk.char_end
            )
            assert parsed.text[chunk.char_start : chunk.char_end] == chunk.text

        indexed = index_sync._to_indexed(document, chunks)
        assert [item.source_locations for item in indexed] == [
            chunk.source_locations for chunk in chunks
        ]

        foreign_tenant_id = uuid.uuid4()
        assert (
            await repositories.DocumentRepository(session, foreign_tenant_id).get(document_id)
            is None
        )
        foreign_chunk_repo = repositories.ChunkRepository(session, foreign_tenant_id)
        assert await foreign_chunk_repo.get(chunks[0].id) is None
