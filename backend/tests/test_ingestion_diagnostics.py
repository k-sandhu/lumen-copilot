"""Extraction measurements and native-table inspection (spec 0016 / #628)."""

from __future__ import annotations

import io
from importlib import import_module

import pytest

from tests import test_ingestion_task as task_fixtures
from tests.test_ingestion_locations import _make_pdf_pages_with_blank_middle
from tests.test_ingestion_task import (
    _FakeGateway,
    _FakeObjectStore,
    _seed_document,
    _settings,
)

sqlite_engine = task_fixtures.sqlite_engine
_offline_index_store = task_fixtures._offline_index_store


def _diagnostics(data: bytes, mime_type: str):
    parsed = import_module("app.ingestion.parsers").parse_document_with_locations(
        data, mime_type=mime_type
    )
    diagnostics_module = import_module("app.ingestion.diagnostics")
    return diagnostics_module.build_extraction_diagnostics(data, mime_type=mime_type, parsed=parsed)


def _make_docx_table() -> bytes:
    import docx

    document = docx.Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Revenue USD"
    table.cell(1, 0).text = "West"
    table.cell(1, 1).text = "Twelve"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _make_docx_split_word_table() -> bytes:
    import docx

    document = docx.Document()
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.add_run("Rev")
    paragraph.add_run("enue USD")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _make_pptx_table() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(5), Inches(2)).table
    table.cell(0, 0).text = "Division"
    table.cell(0, 1).text = "Margin CAD"
    table.cell(1, 0).text = "East"
    table.cell(1, 1).text = "Twenty"
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def _make_sparse_formula_xlsx() -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sparse North"
    sheet["A1"] = "Product"
    sheet["C1"] = "Units"
    sheet["A4"] = "Widget"
    sheet["D7"] = "=SUM(C1:C6)"
    workbook.create_sheet("Empty Region")
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_pdf_reports_blank_native_page_and_unknown_table_coverage() -> None:
    diagnostics = _diagnostics(_make_pdf_pages_with_blank_middle(), mime_type="application/pdf")

    assert diagnostics.source_part_kind == "page"
    assert diagnostics.total_parts == 3
    assert diagnostics.parts_with_text == 2
    assert list(diagnostics.blank_parts) == [2]
    assert diagnostics.table_probe == "unavailable"
    assert diagnostics.table_regions is None
    assert diagnostics.table_cells is None
    assert diagnostics.missing_table_cells is None
    assert any(
        "table" in warning.lower() and "unknown" in warning.lower()
        for warning in diagnostics.warnings
    )


def test_plain_text_counts_replacement_and_suspicious_controls_without_table_warning() -> None:
    data = "A�\x01\t\n\r\x7fB".encode()
    diagnostics = _diagnostics(data, mime_type="text/plain")
    payload = diagnostics.to_dict()

    assert diagnostics.character_count == len("A�\x01\t\n\r\x7fB")
    assert diagnostics.replacement_characters == 1
    assert diagnostics.suspicious_controls == 2
    assert diagnostics.source_part_kind is None
    assert diagnostics.total_parts is None
    assert list(diagnostics.blank_parts) == []
    assert diagnostics.table_probe == "unavailable"
    assert diagnostics.table_regions is None
    assert diagnostics.table_cells is None
    assert diagnostics.missing_table_cells is None
    assert any("replacement" in warning.lower() for warning in diagnostics.warnings)
    assert any("control" in warning.lower() for warning in diagnostics.warnings)
    assert not any("table" in warning.lower() for warning in diagnostics.warnings)
    assert set(payload) == {
        "character_count",
        "replacement_characters",
        "suspicious_controls",
        "source_part_kind",
        "total_parts",
        "parts_with_text",
        "blank_parts",
        "table_probe",
        "table_regions",
        "table_cells",
        "missing_table_cells",
        "warnings",
    }


def test_empty_plain_text_is_distinguished_from_extracted_text() -> None:
    empty = _diagnostics(b"", mime_type="text/plain")
    ordinary = _diagnostics(b"ordinary text", mime_type="text/plain")

    assert empty.character_count == 0
    assert ordinary.character_count == len("ordinary text")
    assert empty.replacement_characters == ordinary.replacement_characters == 0
    assert empty.suspicious_controls == ordinary.suspicious_controls == 0
    assert empty.to_dict() != ordinary.to_dict()


@pytest.mark.parametrize(
    ("mime_type", "data", "cell_values"),
    [
        (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            _make_docx_table(),
            ["Region", "Revenue USD", "West", "Twelve"],
        ),
        (
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            _make_pptx_table(),
            ["Division", "Margin CAD", "East", "Twenty"],
        ),
    ],
    ids=["docx", "pptx"],
)
def test_native_tables_measure_omitted_cells_without_storing_cell_text(
    mime_type: str, data: bytes, cell_values: list[str]
) -> None:
    parsed = import_module("app.ingestion.parsers").parse_document_with_locations(
        data, mime_type=mime_type
    )
    diagnostics = import_module("app.ingestion.diagnostics").build_extraction_diagnostics(
        data, mime_type=mime_type, parsed=parsed
    )
    payload = diagnostics.to_dict()

    assert diagnostics.table_probe == "native_tables"
    assert diagnostics.table_regions == 1
    assert diagnostics.table_cells == 4
    assert diagnostics.missing_table_cells == len(cell_values)
    assert all(value not in repr(payload) for value in cell_values)


def test_docx_split_runs_match_contiguous_parsed_cell_text() -> None:
    from app.domain.ingestion import ParsedDocument

    data = _make_docx_split_word_table()
    diagnostics = import_module("app.ingestion.diagnostics").build_extraction_diagnostics(
        data,
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        parsed=ParsedDocument("Revenue USD"),
    )

    assert diagnostics.table_cells == 1
    assert diagnostics.missing_table_cells == 0


def test_sparse_workbook_counts_nonempty_cells_and_formula_without_cached_value() -> None:
    diagnostics = _diagnostics(
        _make_sparse_formula_xlsx(),
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    assert diagnostics.source_part_kind == "sheet"
    assert diagnostics.total_parts == 2
    assert diagnostics.parts_with_text == 1
    assert list(diagnostics.blank_parts) == [2]
    assert diagnostics.table_probe == "sheet_cells"
    assert diagnostics.table_regions == 1
    assert diagnostics.table_cells == 4
    assert diagnostics.missing_table_cells == 1


@pytest.mark.parametrize(
    "mime_type",
    [
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ],
    ids=["docx", "pptx"],
)
def test_corrupt_ooxml_probe_raises_typed_parse_error(mime_type: str) -> None:
    from app.domain.ingestion import ParsedDocument
    from app.ingestion.parsers import DocumentParseError

    diagnostics_module = import_module("app.ingestion.diagnostics")
    with pytest.raises(DocumentParseError):
        diagnostics_module.build_extraction_diagnostics(
            b"not a valid OOXML container",
            mime_type=mime_type,
            parsed=ParsedDocument("already extracted native text"),
        )


@pytest.mark.asyncio
async def test_ingestion_persists_diagnostics_before_embedding_failure(
    sqlite_engine: None,
) -> None:
    db_session = import_module("app.db.session")
    repositories = import_module("app.db.repositories")
    ingest = import_module("app.tasks.ingest")
    from app.tasks.ingest import IngestionError

    tenant_id, document_id = await _seed_document(mime_type="text/plain", key="diagnostics")
    store = _FakeObjectStore()
    gateway = _FakeGateway(fail=True)
    body = b"Native text survives model failure."
    async with db_session.session_scope() as session:
        document = await repositories.DocumentRepository(session, tenant_id).get(document_id)
        assert document is not None
        store.put(str(tenant_id), document.storage_key, body)

    with pytest.raises(IngestionError):
        await ingest.ingest_document_async(
            tenant_id,
            document_id,
            settings=_settings(),
            object_store=store,
            gateway=gateway,
        )

    async with db_session.session_scope() as session:
        document = await repositories.DocumentRepository(session, tenant_id).get(document_id)
    assert document is not None
    stored = document.ingestion_metadata["extraction_diagnostics"]
    expected = _diagnostics(body, "text/plain").to_dict()
    assert stored == expected


@pytest.mark.asyncio
async def test_embedding_failure_keeps_prior_extraction_and_chunks_but_updates_diagnostics(
    sqlite_engine: None,
) -> None:
    db_session = import_module("app.db.session")
    repositories = import_module("app.db.repositories")
    ingest = import_module("app.tasks.ingest")
    from app.tasks.ingest import IngestionError

    tenant_id, document_id = await _seed_document(mime_type="text/plain", key="diagnostics-rerun")
    store = _FakeObjectStore()
    storage_key: str
    async with db_session.session_scope() as session:
        document = await repositories.DocumentRepository(session, tenant_id).get(document_id)
        assert document is not None
        storage_key = document.storage_key

    old_body = ("Prior source with retained citation spans. " * 12).encode()
    store.put(str(tenant_id), storage_key, old_body)
    first = await ingest.ingest_document_async(
        tenant_id,
        document_id,
        settings=_settings(),
        object_store=store,
        gateway=_FakeGateway(),
    )
    assert first.chunk_count > 0
    async with db_session.session_scope() as session:
        document = await repositories.DocumentRepository(session, tenant_id).get(document_id)
        chunks = await repositories.ChunkRepository(session, tenant_id).list_for_document(
            document_id
        )
    assert document is not None and document.source_text is not None
    prior_text = document.source_text
    prior_locations = document.source_locations
    prior_chunks = [(chunk.text, chunk.char_start, chunk.char_end, chunk.ord) for chunk in chunks]
    assert prior_chunks
    assert all(prior_text[item[1] : item[2]] == item[0] for item in prior_chunks)

    new_body = b"Latest extraction is measured before embeddings can fail."
    store.put(str(tenant_id), storage_key, new_body)
    with pytest.raises(IngestionError):
        await ingest.ingest_document_async(
            tenant_id,
            document_id,
            settings=_settings(),
            object_store=store,
            gateway=_FakeGateway(fail=True),
        )

    async with db_session.session_scope() as session:
        updated = await repositories.DocumentRepository(session, tenant_id).get(document_id)
        chunks = await repositories.ChunkRepository(session, tenant_id).list_for_document(
            document_id
        )
    assert updated is not None
    assert updated.source_text == prior_text
    assert updated.source_locations == prior_locations
    current_chunks = [(chunk.text, chunk.char_start, chunk.char_end, chunk.ord) for chunk in chunks]
    assert current_chunks == prior_chunks
    assert all(updated.source_text[item[1] : item[2]] == item[0] for item in current_chunks)
    diagnostics = updated.ingestion_metadata["extraction_diagnostics"]
    assert diagnostics["character_count"] == len(new_body.decode())
    assert diagnostics["character_count"] != len(prior_text)
